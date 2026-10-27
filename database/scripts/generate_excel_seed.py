#!/usr/bin/env python3
"""Generate deterministic PostgreSQL seed data from the approved demo workbook."""

from __future__ import annotations

import hashlib
import json
import re
import uuid
from collections import defaultdict
from datetime import date, datetime, time, timezone, timedelta
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "销售拜访过程记录及分析_近一个月_五人随机分配.xlsx"
OUTPUT = ROOT / "database" / "seeds" / "S003__excel_demo_data.sql"
WORKSPACE_ID = "00000000-0000-0000-0000-000000000001"
FORM_VERSION_ID = "31000000-0000-0000-0000-000000000002"
QUADRANT_RULE_ID = "40000000-0000-0000-0000-000000000001"
CN_TZ = timezone(timedelta(hours=8))

USER_IDS = {
    "韩启微": "02000000-0000-0000-0000-000000000001",
    "郭强健": "02000000-0000-0000-0000-000000000002",
    "蒋磊": "02000000-0000-0000-0000-000000000003",
    "刘志德": "02000000-0000-0000-0000-000000000004",
    "王新源": "02000000-0000-0000-0000-000000000005",
    "周莹": "02000000-0000-0000-0000-000000000006",
}

TEAM_BY_USER = {
    "郭强健": "01000000-0000-0000-0000-000000000002",
    "刘志德": "01000000-0000-0000-0000-000000000002",
    "王新源": "01000000-0000-0000-0000-000000000002",
    "蒋磊": "01000000-0000-0000-0000-000000000003",
    "周莹": "01000000-0000-0000-0000-000000000003",
}

HEADERS = [
    "客户名称", "合作伙伴", "商机名称", "预计金额（元）", "客户/线索来源", "拜访对象类别",
    "客户职位", "客户名字", "记录人", "拜访及沟通日期", "拜访及沟通方式",
    "拜访地点", "沟通时长", "是/否达成预期", "跟进记录", "下一步",
]

IDENTITY_HEADERS = [header for header in HEADERS if header != "预计金额（元）"]

FIELD_KEYS = [
    "customer_name", "partner_name", "opportunity_name", "estimated_amount", "lead_source", "contact_category",
    "contact_title", "contact_name", "recorder_user_id", "interaction_at", "interaction_mode",
    "visit_location", "duration_minutes", "expectation_met", "follow_up_record", "next_action",
]

FIELD_IDS = {
    "customer_name": "20000000-0000-0000-0000-000000000001",
    "partner_name": "20000000-0000-0000-0000-000000000002",
    "opportunity_name": "20000000-0000-0000-0000-000000000003",
    "lead_source": "20000000-0000-0000-0000-000000000004",
    "contact_category": "20000000-0000-0000-0000-000000000005",
    "contact_title": "20000000-0000-0000-0000-000000000006",
    "contact_name": "20000000-0000-0000-0000-000000000007",
    "recorder_user_id": "20000000-0000-0000-0000-000000000008",
    "interaction_at": "20000000-0000-0000-0000-000000000009",
    "interaction_mode": "20000000-0000-0000-0000-000000000010",
    "visit_location": "20000000-0000-0000-0000-000000000011",
    "duration_minutes": "20000000-0000-0000-0000-000000000012",
    "expectation_met": "20000000-0000-0000-0000-000000000013",
    "follow_up_record": "20000000-0000-0000-0000-000000000014",
    "next_action": "20000000-0000-0000-0000-000000000015",
    "estimated_amount": "20000000-0000-0000-0000-000000000016",
}

EXPECTATION_SCORE = {
    "超出100%": 100, "达成100%": 100, "达成50-100%": 75,
    "达成30-50%": 40, "达成10-30%": 20, "未达成": 0,
}

EXPECTATION_CODE = {
    "超出100%": "exceeded", "达成100%": "met", "达成50-100%": "met_50_100",
    "达成30-50%": "met_30_50", "达成10-30%": "met_10_30", "未达成": "not_met",
}

MODE_CODE = {
    "线下会议": "offline_meeting", "线上会议": "online_meeting", "线上": "online_meeting",
    "电话/语音": "phone_voice", "饭局聚会": "social_meal", "饭局/聚会": "social_meal",
}


def stable_uuid(kind: str, value: str) -> str:
    return str(uuid.uuid5(uuid.NAMESPACE_URL, f"sales-management/{kind}/{value}"))


def sql(value) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        return str(value)
    return "'" + str(value).replace("'", "''") + "'"


def json_sql(value) -> str:
    return sql(json.dumps(value, ensure_ascii=False, separators=(",", ":"))) + "::jsonb"


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def interaction_at(value) -> datetime:
    if isinstance(value, datetime):
        return value.replace(hour=9, minute=0, second=0, microsecond=0, tzinfo=CN_TZ)
    if isinstance(value, date):
        return datetime.combine(value, time(9, 0), tzinfo=CN_TZ)
    return datetime.strptime(text(value)[:10], "%Y-%m-%d").replace(hour=9, tzinfo=CN_TZ)


def amount(value) -> int:
    if isinstance(value, (int, float)):
        result = int(round(value))
    else:
        raw = text(value).replace(",", "").replace("¥", "").replace("￥", "")
        multiplier = 10000 if "万" in raw else 1
        raw = raw.replace("万元", "").replace("万", "").replace("元", "")
        result = int(round(float(raw) * multiplier)) if raw else 0
    if result <= 0:
        raise ValueError("预计金额（元）是必填项且必须大于 0")
    return result


def split_contacts(value: str) -> list[str]:
    return [part.strip() for part in re.split(r"[，、,;/；]", value or "") if part.strip()]


def potential_score(records: list[dict]) -> int:
    latest = records[0]
    score = 35
    if latest["商机名称"]:
        score += 18
    category = latest["拜访对象类别"]
    score += 18 if category == "最终客户" else 14 if category in {"白金伙伴", "金牌伙伴"} else 9 if category else 0
    if latest["客户/线索来源"]:
        score += 8
    if EXPECTATION_SCORE.get(latest["是/否达成预期"], 50) >= 75:
        score += 10
    if len(records) >= 2:
        score += 5
    return min(96, score)


def relationship_score(records: list[dict]) -> int:
    latest = records[0]
    contacts = {name for row in records for name in split_contacts(row["客户名字"])}
    score = 18 + min(42, len(records) * 14)
    if contacts:
        score += 8
    if re.search(r"CXO|VP|总监|部门总|主管|经理", latest["客户职位"] or ""):
        score += 10
    if EXPECTATION_SCORE.get(latest["是/否达成预期"], 50) >= 75:
        score += 10
    return min(96, score)


def read_rows() -> list[dict]:
    workbook = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
    sheet = workbook["数据表"]
    rows = []
    header = [text(cell) for cell in next(sheet.iter_rows(values_only=True))]
    positions = {name: header.index(name) for name in HEADERS}
    for source_row_no, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        row = {
            name: amount(values[index]) if name == "预计金额（元）" else text(values[index])
            for name, index in positions.items()
        }
        if not row["客户名称"]:
            continue
        row["_source_row_no"] = source_row_no
        row["_interaction_at"] = interaction_at(values[positions["拜访及沟通日期"]])
        rows.append(row)
    return rows


def insert_many(lines: list[str], table: str, columns: list[str], values: list[list], conflict: str) -> None:
    if not values:
        return
    lines.append(f"INSERT INTO {table} ({', '.join(columns)}) VALUES")
    lines.append(",\n".join("  (" + ", ".join(row) + ")" for row in values))
    lines.append(conflict + ";\n")


def build() -> None:
    rows = read_rows()
    by_customer = defaultdict(list)
    for row in rows:
        by_customer[row["客户名称"]].append(row)
    for records in by_customer.values():
        records.sort(key=lambda item: (item["_interaction_at"], item["_source_row_no"]), reverse=True)

    customer_ids = {name: stable_uuid("customer", name) for name in by_customer}
    opportunity_ids = {}
    contact_ids = {}
    lines = [
        "\\set ON_ERROR_STOP on",
        "BEGIN;",
        "SET TIME ZONE 'UTC';",
        "SET LOCAL row_security = off;",
        "",
        f"-- Generated from {SOURCE.name}; {len(rows)} visits and {len(by_customer)} customers.",
        "-- Regenerate with database/scripts/generate_excel_seed.py when the approved workbook changes.\n",
    ]

    customer_values = []
    quadrant_values = []
    for name, records in sorted(by_customer.items()):
        latest = records[0]
        owner_name = latest["记录人"]
        potential = potential_score(records)
        relationship = relationship_score(records)
        quadrant = "customer_asset" if potential >= 70 and relationship >= 70 else "main_attack" if potential >= 70 else "customer_resource" if relationship >= 70 else "order_driven"
        customer_values.append([
            sql(customer_ids[name]), sql(WORKSPACE_ID), sql(f"excel-customer-{customer_ids[name]}"), sql(name), sql(name.lower()),
            "NULL", sql(latest["拜访对象类别"] or None), sql(latest["客户/线索来源"] or None), sql("active"),
            sql("A" if potential >= 80 else "B"), sql(USER_IDS[owner_name]), sql(TEAM_BY_USER[owner_name]),
            sql(latest["合作伙伴"] or None), "NULL", sql("excel_import"), sql(USER_IDS[owner_name]),
            json_sql({"sourceWorkbook": SOURCE.name, "visitCount": len(records)}),
        ])
        quadrant_values.append([
            sql(stable_uuid("quadrant", name)), sql(WORKSPACE_ID), sql(customer_ids[name]), str(potential), str(relationship), sql(quadrant),
            sql(QUADRANT_RULE_ID), json_sql({"visitCount": len(records), "latestVisitDate": latest["_interaction_at"].date().isoformat()}),
            json_sql([{"source": "excel_import", "workbook": SOURCE.name}]), sql("2026-09-02 23:59:59+08"),
        ])

    insert_many(lines, "crm.customer", [
        "id", "workspace_id", "external_customer_id", "name", "normalized_name", "industry_code", "customer_type_code", "source_code",
        "lifecycle_status", "level_code", "owner_user_ref_id", "owner_team_id", "primary_partner_name", "demand_summary", "data_source",
        "created_by_user_ref_id", "import_meta",
    ], customer_values, "ON CONFLICT (id) DO UPDATE SET owner_user_ref_id = EXCLUDED.owner_user_ref_id, owner_team_id = EXCLUDED.owner_team_id, level_code = EXCLUDED.level_code, import_meta = EXCLUDED.import_meta")

    opportunity_values = []
    for row in rows:
        key = (row["客户名称"], row["商机名称"])
        if not row["商机名称"] or key in opportunity_ids:
            continue
        oid = stable_uuid("opportunity", "|".join(key))
        opportunity_ids[key] = oid
        owner = row["记录人"]
        opportunity_values.append([
            sql(oid), sql(WORKSPACE_ID), sql(customer_ids[row["客户名称"]]), sql(f"excel-opportunity-{oid}"), sql(row["商机名称"]),
            sql("identified"), sql(row["预计金额（元）"]), sql("CNY"), sql(USER_IDS[owner]), sql(TEAM_BY_USER[owner]), sql("open"), sql(row["客户/线索来源"] or None),
            sql(USER_IDS[owner]), json_sql({"sourceWorkbook": SOURCE.name}),
        ])
    insert_many(lines, "crm.opportunity", [
        "id", "workspace_id", "customer_id", "external_opportunity_id", "name", "stage_code", "amount", "currency", "owner_user_ref_id", "owner_team_id", "status",
        "source_code", "created_by_user_ref_id", "import_meta",
    ], opportunity_values, "ON CONFLICT (id) DO UPDATE SET amount = EXCLUDED.amount, currency = EXCLUDED.currency, owner_user_ref_id = EXCLUDED.owner_user_ref_id, owner_team_id = EXCLUDED.owner_team_id, import_meta = EXCLUDED.import_meta")

    contact_values = []
    for customer_name, records in sorted(by_customer.items()):
        seen = set()
        for row in records:
            for contact_name in split_contacts(row["客户名字"]):
                key = (customer_name, contact_name, row["客户职位"])
                if key in seen:
                    continue
                seen.add(key)
                cid = stable_uuid("contact", "|".join(key))
                contact_ids[key] = cid
                contact_values.append([
                    sql(cid), sql(WORKSPACE_ID), sql(customer_ids[customer_name]), sql(f"excel-contact-{cid}"), sql(contact_name),
                    sql(row["客户职位"] or None), sql(row["拜访对象类别"] or None), "TRUE" if len(seen) == 1 else "FALSE",
                    sql(USER_IDS[row["记录人"]]), json_sql({"sourceWorkbook": SOURCE.name}),
                ])
    insert_many(lines, "crm.contact", [
        "id", "workspace_id", "customer_id", "external_contact_id", "name", "title", "contact_category_code", "is_primary",
        "created_by_user_ref_id", "import_meta",
    ], contact_values, "ON CONFLICT (id) DO UPDATE SET title = EXCLUDED.title, contact_category_code = EXCLUDED.contact_category_code, import_meta = EXCLUDED.import_meta")

    visit_values = []
    visit_contact_values = []
    field_values = []
    for row in rows:
        raw_identity = "|".join(text(row[name]) for name in IDENTITY_HEADERS) + f"|row:{row['_source_row_no']}"
        visit_id = stable_uuid("visit", raw_identity)
        fingerprint = hashlib.sha256(raw_identity.encode("utf-8")).hexdigest()
        opportunity_id = opportunity_ids.get((row["客户名称"], row["商机名称"]))
        owner = row["记录人"]
        occurred = row["_interaction_at"].isoformat()
        visit_values.append([
            sql(visit_id), sql(WORKSPACE_ID), sql(customer_ids[row["客户名称"]]), sql(opportunity_id), sql(USER_IDS[owner]), sql(TEAM_BY_USER[owner]),
            sql(FORM_VERSION_ID), sql("archived"), sql(occurred), sql(MODE_CODE.get(row["拜访及沟通方式"], row["拜访及沟通方式"] or None)),
            sql(row["拜访地点"] or None), "NULL", sql(EXPECTATION_CODE.get(row["是/否达成预期"], row["是/否达成预期"] or None)),
            sql(row["跟进记录"] or None), sql(row["下一步"] or None), sql(row["合作伙伴"] or None), sql(row["客户/线索来源"] or None),
            sql(row["拜访对象类别"] or None), sql(row["客户职位"] or None), sql(row["客户名字"] or None), sql(fingerprint),
            sql(USER_IDS[owner]), sql(occurred), sql(occurred), json_sql({"sourceWorkbook": SOURCE.name, "sourceRow": row["_source_row_no"], "timePrecision": "date"}),
        ])
        for contact_name in split_contacts(row["客户名字"]):
            ckey = (row["客户名称"], contact_name, row["客户职位"])
            if ckey in contact_ids:
                visit_contact_values.append([sql(visit_id), sql(contact_ids[ckey]), sql(WORKSPACE_ID), sql("attendee")])
        source_values = [row[name] if name == "预计金额（元）" else text(row[name]) for name in HEADERS]
        source_values[8] = owner
        source_values[9] = occurred
        for index, (field_key, value) in enumerate(zip(FIELD_KEYS, source_values)):
            field_id = FIELD_IDS[field_key]
            field_values.append([
                sql(visit_id), sql(WORKSPACE_ID), sql(field_id), json_sql(value), sql(value), sql("import"), "NULL", "TRUE",
                json_sql([{"source": "workbook", "file": SOURCE.name, "row": row["_source_row_no"], "column": HEADERS[index]}]),
            ])

    insert_many(lines, "activity.visit", [
        "id", "workspace_id", "customer_id", "opportunity_id", "recorder_user_ref_id", "recorder_team_id", "form_version_id", "status",
        "interaction_at", "interaction_mode_code", "visit_location", "duration_minutes", "expectation_code", "follow_up_record", "next_action",
        "partner_name_snapshot", "lead_source_snapshot", "contact_category_snapshot", "contact_title_snapshot", "contact_name_snapshot",
        "idempotency_fingerprint", "confirmed_by_user_ref_id", "confirmed_at", "archived_at", "import_meta",
    ], visit_values, "ON CONFLICT (id) DO UPDATE SET form_version_id = EXCLUDED.form_version_id, follow_up_record = EXCLUDED.follow_up_record, next_action = EXCLUDED.next_action, import_meta = EXCLUDED.import_meta")

    insert_many(lines, "activity.visit_contact", ["visit_id", "contact_id", "workspace_id", "participation_role"], visit_contact_values,
                "ON CONFLICT (visit_id, contact_id) DO UPDATE SET participation_role = EXCLUDED.participation_role")
    insert_many(lines, "activity.visit_field_value", [
        "visit_id", "workspace_id", "field_definition_id", "value_json", "value_text", "source_type", "confidence", "is_confirmed", "evidence",
    ], field_values, "ON CONFLICT (visit_id, field_definition_id) DO UPDATE SET value_json = EXCLUDED.value_json, value_text = EXCLUDED.value_text, evidence = EXCLUDED.evidence")

    insert_many(lines, "insight.quadrant_score", [
        "id", "workspace_id", "customer_id", "potential_score", "relationship_score", "quadrant_code", "rule_set_id", "input_snapshot", "evidence", "calculated_at",
    ], quadrant_values, "ON CONFLICT (id) DO UPDATE SET potential_score = EXCLUDED.potential_score, relationship_score = EXCLUDED.relationship_score, quadrant_code = EXCLUDED.quadrant_code, input_snapshot = EXCLUDED.input_snapshot")

    lines.extend([
        "COMMIT;",
        "",
        f"-- Expected: {len(by_customer)} customers, {len(opportunity_values)} opportunities, {len(contact_values)} contacts, {len(rows)} visits, {len(field_values)} visit field values.",
    ])
    OUTPUT.write_text("\n".join(lines), encoding="utf-8")
    print(json.dumps({
        "source": str(SOURCE), "output": str(OUTPUT), "customers": len(by_customer),
        "opportunities": len(opportunity_values), "contacts": len(contact_values),
        "visits": len(rows), "visit_field_values": len(field_values),
    }, ensure_ascii=False))


if __name__ == "__main__":
    build()
